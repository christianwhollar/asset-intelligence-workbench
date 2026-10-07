"""Bounded spreadsheet ingestion with full-snapshot validation and two-person publication."""

import calendar
import csv
import hashlib
import io
import json
import re
import uuid
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from .db import Problem, audit, current, now, packed

HEADERS = {
    "Assets": [
        "asset_id",
        "aliases",
        "name",
        "sector",
        "market",
        "area_sf",
        "annual_debt_service",
        "replacement_reserve",
        "cap_rate",
    ],
    "Leases": [
        "lease_id",
        "asset_ref",
        "tenant_name",
        "area_sf",
        "rent",
        "rent_frequency",
        "start_date",
        "end_date",
    ],
    "Operating": ["asset_ref", "month", "account", "amount"],
}
INCOME = {"rent", "other_income"}
EXPENSE = {"tax", "insurance", "utilities", "repairs", "management"}
EXCLUDED = {"capex", "debt_service", "depreciation"}
LIMIT = 2_000_000


def months(as_of):
    d = date.fromisoformat(as_of)
    return [
        f"{(d.year * 12 + d.month - 1 - i) // 12:04d}-{(d.year * 12 + d.month - 1 - i) % 12 + 1:02d}"
        for i in reversed(range(12))
    ]


def text(value, limit=160):
    result = str(value or "").strip()
    if not result or len(result) > limit or any(ord(c) < 32 for c in result):
        raise ValueError("Missing, oversized, or invalid text field")
    return result


def number(value, money=False):
    raw = str(value).strip().replace("$", "").replace(",", "")
    try:
        n = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError("Expected a number") from exc
    if not n.is_finite() or n < 0 or n > Decimal("100000000000"):
        raise ValueError("Numeric value outside supported range")
    scaled = n * (100 if money else 1)
    if scaled != scaled.to_integral_value():
        raise ValueError("Money requires whole cents; area requires whole square feet")
    return int(scaled)


def day(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return date.fromisoformat(str(value)).isoformat()


def read_tables(data, filename):
    if len(data) > LIMIT:
        raise Problem("Upload exceeds 2 MB")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        if sum(i.file_size for i in archive.infolist()) > 8_000_000 or len(archive.infolist()) > 100:
            raise Problem("Expanded archive exceeds limits")
        tables = {}
        if filename.lower().endswith(".xlsx"):
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            if set(book.sheetnames) != set(HEADERS):
                raise Problem("Workbook must contain exactly Assets, Leases, and Operating sheets")
            for name in HEADERS:
                rows = []
                for row in book[name].iter_rows():
                    if len(rows) > 6000 or len(row) > 16:
                        raise Problem("Sheet exceeds 6,000 rows or 16 columns")
                    if any(c.data_type == "f" for c in row):
                        raise Problem("Formula cells are not accepted. Export values before importing.")
                    rows.append([c.value for c in row])
                tables[name] = rows
            book.close()
        elif filename.lower().endswith(".zip"):
            names = {name + ".csv" for name in HEADERS}
            if set(archive.namelist()) != names or len(archive.infolist()) != 3:
                raise Problem("CSV bundle must contain only Assets.csv, Leases.csv, and Operating.csv")
            for name in HEADERS:
                rows = list(csv.reader(io.StringIO(archive.read(name + ".csv").decode("utf-8-sig"))))
                if len(rows) > 6001 or any(len(r) > 16 for r in rows):
                    raise Problem("CSV exceeds row or column limits")
                tables[name] = rows
        else:
            raise Problem("Upload an .xlsx workbook or .zip containing the three CSV files")
        out = {}
        for name, rows in tables.items():
            if not rows or rows[0] != HEADERS[name]:
                raise Problem(f"{name}: headers must match the downloadable template exactly")
            out[name] = []
            for i, row in enumerate(rows[1:], 2):
                if not any(v is not None and v != "" for v in row):
                    continue
                if len(row) != len(HEADERS[name]):
                    raise Problem(f"{name} row {i}: wrong column count")
                out[name].append(dict(zip(HEADERS[name], row), source_row=i))
        return out
    except Problem:
        raise
    except Exception as exc:
        raise Problem("Cannot read workbook. Check the file format, encoding, and template.") from exc


def validate(tables, as_of):
    try:
        d = date.fromisoformat(as_of)
        if d.day != calendar.monthrange(d.year, d.month)[1] or d > date.today():
            raise ValueError()
    except ValueError as exc:
        raise Problem("As-of date must be the last day of a completed reporting month") from exc
    errors, aliases, assets, leases, operating = [], {}, [], [], []

    def error(sheet, row, message):
        errors.append({"sheet": sheet, "row": row, "message": str(message)})

    for row in tables["Assets"]:
        try:
            aid = text(row["asset_id"], 40)
            if not re.fullmatch(r"[A-Za-z0-9_-]+", aid):
                raise ValueError("Asset ID must use letters, numbers, underscores, or hyphens")
            names = {
                aid.casefold(),
                *[x.strip().casefold() for x in str(row["aliases"] or "").split("|") if x.strip()],
            }
            if any(n in aliases for n in names):
                raise ValueError("Duplicate asset ID or ambiguous alias")
            area = number(row["area_sf"])
            cap = float(str(row["cap_rate"]).strip("%")) / (100 if "%" in str(row["cap_rate"]) else 1)
            if area <= 0 or not 0.005 <= cap <= 0.3:
                raise ValueError("Area must be positive and cap rate between 0.5% and 30%")
            asset = {
                "id": aid,
                "name": text(row["name"]),
                "sector": text(row["sector"]),
                "market": text(row["market"]),
                "area_sf": area,
                "debt_cents": number(row["annual_debt_service"], True),
                "reserve_cents": number(row["replacement_reserve"], True),
                "cap_rate": cap,
                "source_row": row["source_row"],
            }
            assets.append(asset)
            aliases.update({n: aid for n in names})
        except (ValueError, TypeError, OverflowError) as exc:
            error("Assets", row["source_row"], exc)
    ids, used_ops = set(), set()

    def asset_ref(row):
        ref = str(row["asset_ref"]).strip().casefold()
        if ref not in aliases:
            raise ValueError("Unknown asset reference; add an explicit alias on Assets")
        return aliases[ref]

    for row in tables["Leases"]:
        try:
            lid = text(row["lease_id"], 60)
            if lid in ids:
                raise ValueError("Duplicate lease ID")
            aid = asset_ref(row)
            start, end = day(row["start_date"]), day(row["end_date"])
            if end < start:
                raise ValueError("Lease end precedes start")
            frequency = str(row["rent_frequency"]).lower().strip()
            if frequency not in {"annual", "monthly"}:
                raise ValueError("rent_frequency must be annual or monthly")
            area = number(row["area_sf"])
            if area == 0:
                raise ValueError("Lease area must be positive")
            leases.append(
                {
                    "id": lid,
                    "asset": aid,
                    "tenant_name": text(row["tenant_name"]),
                    "area_sf": area,
                    "rent_cents": number(row["rent"], True) * (12 if frequency == "monthly" else 1),
                    "start_date": start,
                    "end_date": end,
                    "source_row": row["source_row"],
                }
            )
            ids.add(lid)
        except (ValueError, TypeError) as exc:
            error("Leases", row["source_row"], exc)
    window = months(as_of)
    for row in tables["Operating"]:
        try:
            aid, month, account = asset_ref(row), str(row["month"]), str(row["account"]).strip().lower()
            if month not in window:
                raise ValueError("Month outside the trailing 12 completed months")
            if account not in INCOME | EXPENSE | EXCLUDED:
                raise ValueError("Unknown account; use an explicit supported account mapping")
            key = (aid, month, account)
            if key in used_ops:
                raise ValueError("Duplicate asset/month/account")
            operating.append(
                {
                    "asset": aid,
                    "month": month,
                    "account": account,
                    "amount_cents": number(row["amount"], True),
                    "source_row": row["source_row"],
                }
            )
            used_ops.add(key)
        except (ValueError, TypeError) as exc:
            error("Operating", row["source_row"], exc)
    if not assets or len(assets) > 100:
        error("Assets", 1, "A snapshot must contain 1–100 assets")
    for a in assets:
        active = [
            lease
            for lease in leases
            if lease["asset"] == a["id"] and lease["start_date"] <= as_of <= lease["end_date"]
        ]
        if sum(lease["area_sf"] for lease in active) > a["area_sf"]:
            error("Leases", 1, f"{a['id']}: active leased area exceeds property area")
        for month in window:
            missing = (INCOME | EXPENSE) - {x[2] for x in used_ops if x[:2] == (a["id"], month)}
            if missing:
                error(
                    "Operating",
                    1,
                    f"{a['id']} {month}: missing accounts {', '.join(sorted(missing))}; enter explicit zeros",
                )
    payload = {"assets": assets, "leases": leases, "operating": operating}
    report = {
        "valid": not errors,
        "errors": errors[:100],
        "error_count": len(errors),
        "counts": {k: len(v) for k, v in payload.items()},
        "as_of": as_of,
        "normalization": "Explicit asset aliases; monthly rent annualized; USD values stored in cents.",
        "excluded_accounts": sorted(EXCLUDED),
    }
    return payload, report


def stage(conn, who, data, filename, as_of, parsed=None):
    checksum = hashlib.sha256(data).hexdigest()
    existing = conn.execute(
        "SELECT id,report,status FROM datasets WHERE tenant=? AND checksum=? AND as_of=?",
        (who.tenant, checksum, as_of),
    ).fetchone()
    if existing:
        return {
            "id": existing["id"],
            "report": json.loads(existing["report"]),
            "status": existing["status"],
            "reused": True,
        }
    payload, report = parsed if parsed is not None else validate(read_tables(data, filename), as_of)
    ident = uuid.uuid4().hex
    conn.execute(
        "INSERT INTO datasets VALUES(?,?,?,?,?,?,?,?,?)",
        (ident, who.tenant, as_of, checksum, who.user, now(), "staged", packed(report), packed(payload)),
    )
    audit(conn, who, "import.staged", {"id": ident, "checksum": checksum, "valid": report["valid"]})
    return {"id": ident, "report": report, "status": "staged", "reused": False}


def publish(conn, who, ident, expected):
    if who.role != "reviewer":
        raise Problem("An independent reviewer must publish the snapshot", 403)
    row = conn.execute("SELECT * FROM datasets WHERE id=? AND tenant=?", (ident, who.tenant)).fetchone()
    if not row:
        raise Problem("Import not found", 404)
    if row["actor"] == who.user:
        raise Problem("The importer cannot review their own snapshot", 403)
    active = conn.execute("SELECT dataset FROM portfolios WHERE tenant=?", (who.tenant,)).fetchone()
    active_id = active[0] if active else None
    if active_id == ident:
        return {"id": ident, "published": True, "reused": True}
    if expected != active_id:
        raise Problem("Portfolio changed since preview; refresh before publishing", 409)
    if not json.loads(row["report"])["valid"]:
        raise Problem("Fix every validation error before publication")
    # Previously accepted versions can be restored explicitly with optimistic concurrency.
    if row["status"] != "published":
        data = json.loads(row["payload"])
        for table in ("assets", "leases", "operating"):
            for record in data[table]:
                columns = ["dataset"] + list(record)
                conn.execute(
                    f"INSERT INTO {table}({','.join(columns)}) VALUES({','.join('?' for _ in columns)})",
                    [ident] + list(record.values()),
                )
        conn.execute("UPDATE datasets SET status='published' WHERE id=?", (ident,))
    conn.execute(
        "INSERT INTO portfolios VALUES(?,?) ON CONFLICT(tenant) DO UPDATE SET dataset=excluded.dataset",
        (who.tenant, ident),
    )
    audit(conn, who, "import.published", {"id": ident, "previous": active_id})
    return {"id": ident, "published": True, "reused": False}


def imports(conn, who):
    active = None
    try:
        active = current(conn, who.tenant)["id"]
    except Problem:
        pass
    rows = conn.execute(
        "SELECT id,as_of,actor,created,status,report FROM datasets WHERE tenant=? ORDER BY created DESC LIMIT 30",
        (who.tenant,),
    ).fetchall()
    return {"current": active, "imports": [dict(r, report=json.loads(r["report"])) for r in rows]}
